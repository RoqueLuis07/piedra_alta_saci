"""Sistema multiusuario: login propio + búsqueda + panel de roles.

Es la app que corre en Railway (ver ``wsgi.py``). Toda la persistencia --
datos, usuarios e incluso las imágenes de las marcas -- vive en un único
Postgres propio (ver ``db.py``). ``requiere_sesion``/``requiere_rol`` (en
``auth.py``) son la barrera de acceso; no hay una segunda capa de RLS como
antes con Supabase, así que estos decoradores son la única barrera real.
"""

from __future__ import annotations

import base64
import io
import os
import re
import tempfile
import zipfile

import click
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from flask import Flask, jsonify, redirect, render_template, request, send_file, session, url_for
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from flask_wtf import CSRFProtect
from flask_wtf.csrf import CSRFError
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import pypdfium2 as pdfium

from marcas.pdf.guia import analizar_venta, completar_guia, completar_venta, extraer_encabezado
from marcas.servidor.auth import (
    cerrar_conexion_actual,
    cerrar_sesion,
    conexion_actual,
    encriptar_contrasena,
    iniciar_sesion,
    perfil_actual,
    requiere_rol,
    requiere_sesion,
    verificar_contrasena,
)
from marcas.servidor.consultas import (
    CAMPOS_MARCA_EDITABLES,
    CAMPOS_OPERACION_EDITABLES,
    CAMPOS_PROPIETARIO_EDITABLES,
    LIMITE_EXPORTAR,
    POR_PAGINA_MARCAS,
    POR_PAGINA_PROPIETARIOS,
    SECCIONES_FILTRO,
    ErrorResolverCambio,
    actualizar_marca_campos,
    actualizar_operacion_campos,
    actualizar_preferencias_columnas,
    actualizar_propietario,
    actualizar_usuario_campos,
    borrador_venta_dominante as obtener_borrador_venta_dominante,
    borrador_venta_pdf as obtener_borrador_venta_pdf,
    buscar_marcas,
    cambio_pendiente_de,
    cambios_pendientes_detalle,
    crear_marca,
    crear_operacion,
    crear_usuario,
    desglose_marcas,
    dict_a_json,
    eliminar_borrador_venta,
    eliminar_filtro,
    estadisticas,
    exportar_marcas,
    exportar_operaciones,
    exportar_todo_para_respaldo,
    ficha_marca,
    guardar_borrador_venta,
    guardar_filtro,
    historial_cambios_resueltos,
    imagen_marca,
    listar_filtros,
    listar_operaciones_paginado,
    listar_propietarios,
    listar_usuarios,
    marcas_a_revisar,
    marcas_de_operacion,
    marcas_de_propietario,
    marcas_por_codigos,
    marcas_por_vencer,
    nombres_usuarios,
    obtener_marca_por_id,
    obtener_operacion_por_id,
    obtener_propietario,
    obtener_usuario_por_email,
    obtener_usuario_por_id,
    ocultar_marca,
    ocultar_operacion,
    ocultar_propietario,
    operaciones_de_propietario,
    preferencias_columnas,
    proponer_cambio,
    ranking_participantes,
    resolver_cambio_pendiente,
    restaurar_marca,
    restaurar_operacion,
    restaurar_propietario,
    resumen_mensual,
    subir_imagen_marca,
    ultimos_asientos,
    url_imagen,
)
from marcas.servidor.db import ejecutar_schema

_CODIGO_VALIDO = re.compile(r"^[A-Za-z0-9._-]+$")


def _error_seguro(mensaje: str, exc: Exception):
    """El detalle real queda en los logs del servidor, nunca en la respuesta:
    puede traer nombres de columnas o de tablas que no hace falta mostrar."""
    print(f"{mensaje}: {exc}")
    return render_template("error_simple.html", titulo="No se pudo completar", mensaje=mensaje), 400


def _registro_no_encontrado(mensaje: str):
    return render_template("error_simple.html", titulo="No encontrado", mensaje=mensaje), 404


def _fecha_sin_tz(valor):
    """openpyxl no admite datetimes con huso horario (falla al guardar el
    .xlsx) -- ``creado_en`` es timestamptz, así que hay que sacarle el tzinfo
    antes de escribirlo en una celda."""
    if isinstance(valor, datetime) and valor.tzinfo is not None:
        return valor.replace(tzinfo=None)
    return valor


def _libro_excel(nombre_hoja: str, encabezados: list[str], filas: list[list]) -> io.BytesIO:
    """Arma un .xlsx prolijo para presentar como informe -- encabezado en
    negrita con fondo verde, congelado arriba, columnas con ancho ajustado
    al contenido y fechas con formato real (no todo como texto plano, como
    quedaba en el CSV viejo)."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre_hoja[:31]
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4D34")
        celda.alignment = Alignment(vertical="center")
    hoja.freeze_panes = "A2"
    for fila in filas:
        hoja.append(fila)
    for fila_celdas in hoja.iter_rows(min_row=2):
        for celda in fila_celdas:
            if isinstance(celda.value, datetime):
                celda.number_format = "DD/MM/YYYY HH:MM"
            elif isinstance(celda.value, date):
                celda.number_format = "DD/MM/YYYY"
    for indice, encabezado in enumerate(encabezados):
        valores = [encabezado] + [fila[indice] for fila in filas]
        ancho = max((len(str(v)) for v in valores if v is not None), default=0)
        hoja.column_dimensions[get_column_letter(indice + 1)].width = min(max(ancho + 2, 10), 42)
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer


def _previsualizar_estampado(conexion, operacion: dict, marcas_elegidas: list[dict], pdf_bytes: bytes, dominante_bytes: bytes | None = None) -> dict:
    """Corre completar_venta/completar_guia con las marcas elegidas hasta el
    momento sobre un PDF y un directorio TEMPORALES -- nunca sobre el PDF
    real que se termina descargando -- y devuelve las páginas que quedaron
    con algo nuevo (como PNG) para el panel izquierdo, más el cupo libre
    restante. Se usa tanto para Venta (Rubro 2 + Anexo) como para Guía (sólo
    Anexo, completar_guia no toca el Rubro 2)."""
    es_venta = operacion.get("tipo_operacion") == "venta"
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        pdf_entrada = tmp_path / "entrada.pdf"
        pdf_entrada.write_bytes(pdf_bytes)

        rutas_imagenes = []
        for marca in marcas_elegidas:
            contenido = imagen_marca(conexion, marca["id"], "png")
            if not contenido:
                continue
            destino = tmp_path / f"{marca['codigo']}.png"
            destino.write_bytes(contenido)
            rutas_imagenes.append(destino)

        dominante_ruta = None
        if dominante_bytes:
            dominante_ruta = tmp_path / "dominante.png"
            dominante_ruta.write_bytes(dominante_bytes)

        salida = tmp_path / "salida.pdf"
        paginas_modificadas: list[int] = []
        error = None
        try:
            if es_venta:
                info = completar_venta(pdf_entrada, dominante_ruta, rutas_imagenes, salida)
            elif rutas_imagenes:
                info = completar_guia(pdf_entrada, rutas_imagenes, salida)
            else:
                info = {"paginas_modificadas": []}
            paginas_modificadas = info.get("paginas_modificadas", [])
        except Exception as exc:
            error = str(exc)
            salida = pdf_entrada  # sin marcas que entren, se previsualiza el PDF tal cual

        paginas_png = []
        fuente = salida if salida.exists() else pdf_entrada
        doc = pdfium.PdfDocument(str(fuente))
        try:
            # Sin nada estampado todavía, se muestra la primera página
            # principal tal cual llegó, para que el panel izquierdo nunca
            # quede vacío antes de la primera marca.
            indices = paginas_modificadas or ([0] if len(doc) else [])
            for i in indices:
                if i >= len(doc):
                    continue
                imagen = doc[i].render(scale=1.3).to_pil()
                buffer = io.BytesIO()
                imagen.save(buffer, format="PNG")
                paginas_png.append(base64.b64encode(buffer.getvalue()).decode("ascii"))
        finally:
            doc.close()

        try:
            cupo = analizar_venta(pdf_entrada)
        except Exception:
            cupo = None

        try:
            encabezado = extraer_encabezado(pdf_entrada)
        except Exception:
            encabezado = None

    return {"paginas": paginas_png, "cupo": cupo, "encabezado": encabezado, "error": error}


def crear_app() -> Flask:
    app = Flask(__name__)
    app.secret_key = os.environ.get("FLASK_SECRET_KEY")
    if not app.secret_key:
        raise RuntimeError("Falta FLASK_SECRET_KEY (variable de entorno) para firmar la sesión.")
    app.config["SESSION_COOKIE_SECURE"] = True
    app.config["SESSION_COOKIE_HTTPONLY"] = True
    app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
    # La sesión ya no depende de un token de un tercero que vence solo --
    # es la propia cookie de Flask, así que le ponemos nosotros un límite
    # (se cierra sola tras 12 horas de inactividad).
    app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=12)
    # Sin este límite, cualquier sesión logueada podría mandar un archivo
    # enorme (imagen de marca o PDF de guía) y agotar memoria/disco del
    # servidor. 20 MB alcanza de sobra para una imagen o un PDF de guía.
    app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024

    # Token oculto obligatorio en todo formulario POST -- protege contra
    # peticiones falsificadas desde otro sitio (CSRF). Queda disponible en
    # los templates como {{ csrf_token() }}, Flask-WTF lo registra solo.
    CSRFProtect(app)

    # Sin esto, el login no tiene freno: alguien podría probar contraseñas
    # sin límite. El almacenamiento en memoria alcanza para un solo proceso;
    # si el día de mañana corremos varios workers en paralelo, pasar a un
    # backend compartido (Redis) para que el límite sea real entre todos.
    limiter = Limiter(get_remote_address, app=app, default_limits=[])

    app.teardown_appcontext(cerrar_conexion_actual)

    @app.cli.command("crear-esquema")
    def crear_esquema_cli():
        """Aplica schema.sql -- idempotente, se puede correr en cada deploy."""
        ejecutar_schema()
        print("Esquema aplicado.")

    @app.cli.command("crear-admin")
    @click.option(
        "--password", "password_opcion", default=None,
        help="Si no se pasa, se pide por consola (visible -- la consola web de Railway "
             "no siempre soporta bien la entrada oculta de contraseñas).",
    )
    def crear_admin_cli(password_opcion):
        """Da de alta la primera cuenta administradora (o cualquier otra,
        a mano) -- para cuando todavía no hay nadie que pueda entrar al Panel."""
        email = input("Email: ").strip().lower()
        nombre = input("Nombre: ").strip()
        contrasena = password_opcion if password_opcion is not None else input("Contraseña (queda visible): ")
        if len(contrasena) < 6:
            print("La contraseña debe tener al menos 6 caracteres.")
            return
        with app.app_context():
            conexion = conexion_actual()
            if obtener_usuario_por_email(conexion, email):
                print(f"Ya existe una cuenta con ese email ({email}).")
                return
            crear_usuario(conexion, email, encriptar_contrasena(contrasena), nombre, "administrador", True)
            conexion.commit()
        print(f"Cuenta administradora creada: {email}")

    @app.get("/login")
    def login():
        if session.get("usuario_id"):
            return redirect(url_for("inicio"))
        return render_template("login.html", error=None)

    @app.post("/login")
    @limiter.limit("8 per minute; 30 per hour")
    def login_post():
        email = request.form.get("email", "").strip().lower()
        contrasena = request.form.get("contrasena", "")
        conexion = conexion_actual()
        usuario = obtener_usuario_por_email(conexion, email)
        if not usuario or not verificar_contrasena(contrasena, usuario["password_hash"]):
            return render_template("login.html", error="Usuario o contraseña incorrectos."), 401
        iniciar_sesion(str(usuario["id"]))
        return redirect(url_for("inicio"))

    @app.post("/logout")
    def logout():
        cerrar_sesion()
        return redirect(url_for("login"))

    @app.get("/cuenta-pendiente")
    def cuenta_pendiente():
        return render_template("cuenta_pendiente.html")

    @app.get("/mi-cuenta")
    @requiere_sesion
    def mi_cuenta():
        return render_template("mi_cuenta.html", activo="mi_cuenta", perfil=perfil_actual(), error=None)

    @app.post("/mi-cuenta")
    @requiere_sesion
    def mi_cuenta_post():
        conexion = conexion_actual()
        perfil = perfil_actual()
        actual = request.form.get("actual", "")
        nueva = request.form.get("nueva", "")
        repetir = request.form.get("repetir", "")
        usuario = obtener_usuario_por_email(conexion, perfil["email"])
        error = None
        if not verificar_contrasena(actual, usuario["password_hash"]):
            error = "La contraseña actual no es correcta."
        elif len(nueva) < 6:
            error = "La contraseña nueva debe tener al menos 6 caracteres."
        elif nueva != repetir:
            error = "Las dos contraseñas nuevas no coinciden."
        if error:
            return render_template("mi_cuenta.html", activo="mi_cuenta", perfil=perfil, error=error), 400
        actualizar_usuario_campos(conexion, perfil["id"], {"password_hash": encriptar_contrasena(nueva)})
        return redirect(url_for("mi_cuenta"))

    @app.errorhandler(429)
    def demasiados_intentos(_exc):
        return render_template(
            "error_simple.html",
            titulo="Demasiados intentos",
            mensaje="Probaste demasiadas veces en poco tiempo. Esperá unos minutos y volvé a intentar.",
        ), 429

    @app.errorhandler(413)
    def archivo_demasiado_grande(_exc):
        return render_template(
            "error_simple.html",
            titulo="El archivo es demasiado grande",
            mensaje="El límite es 20 MB por archivo. Probá con una imagen o un PDF más liviano.",
        ), 413

    @app.errorhandler(CSRFError)
    def token_invalido(_exc):
        return render_template(
            "error_simple.html",
            titulo="La página venció",
            mensaje="El formulario tardó demasiado o se abrió en otra pestaña. Volvé atrás y probá de nuevo.",
        ), 400

    @app.errorhandler(404)
    def no_encontrado(_exc):
        # Cubre tanto una URL que no matchea ninguna ruta (ej. /guias/abc con
        # un id que no es numérico) como un abort(404) explícito -- antes de
        # esto, el primer caso mostraba la página cruda de Flask, en inglés.
        return render_template(
            "error_simple.html",
            titulo="No encontrado",
            mensaje="No existe la página o el registro que buscás. Puede que el enlace esté mal escrito o ya no exista.",
        ), 404

    @app.errorhandler(405)
    def metodo_no_permitido(_exc):
        return render_template(
            "error_simple.html",
            titulo="Esa acción no se puede hacer así",
            mensaje="Probaste acceder directo a una dirección pensada para un formulario. Usá los botones y enlaces del sitio en vez de escribir la URL a mano.",
        ), 405

    @app.context_processor
    def inyectar_pendientes():
        perfil = perfil_actual()
        if not perfil or perfil.get("rol") not in ("administrador", "operador"):
            return {}
        try:
            conexion = conexion_actual()
            with conexion.cursor() as cur:
                cur.execute("SELECT count(*) AS total FROM cambios_pendientes WHERE estado = 'pendiente'")
                n = cur.fetchone()["total"]
        except Exception:
            n = 0
        return {"pendientes_nav": n}

    @app.get("/")
    @requiere_sesion
    def inicio():
        conexion = conexion_actual()
        return render_template(
            "inicio.html",
            activo="inicio",
            perfil=perfil_actual(),
            stats=estadisticas(conexion),
            asientos=ultimos_asientos(conexion),
            a_revisar=marcas_a_revisar(conexion),
            a_vencer=marcas_por_vencer(conexion),
        )

    @app.get("/buscar")
    @requiere_sesion
    def buscar():
        conexion = conexion_actual()
        perfil = perfil_actual()
        texto = request.args.get("q") or None
        estado = request.args.get("estado") or None
        tipo = request.args.get("tipo") or None
        pagina = max(1, request.args.get("pagina", 1, type=int))
        ver_ocultos = perfil["rol"] == "administrador" and request.args.get("ver") == "ocultos"
        resultados, total = buscar_marcas(
            conexion, texto, pagina=pagina, estado=estado, tipo=tipo, incluir_ocultos=ver_ocultos
        )
        for fila in resultados:
            fila["imagen_url"] = url_imagen(fila["id"], fila.get("archivo_png"))
        return render_template(
            "buscar.html",
            activo="buscar",
            resultados=resultados,
            total=total,
            pagina=pagina,
            por_pagina=POR_PAGINA_MARCAS,
            texto=texto or "",
            estado=estado or "",
            tipo=tipo or "",
            ver_ocultos=ver_ocultos,
            perfil=perfil,
            filtros_guardados=listar_filtros(conexion, perfil["id"], "marcas"),
            columnas_visibles=preferencias_columnas(conexion, perfil["id"], "marcas"),
        )

    @app.get("/marcas/<codigo>")
    @requiere_sesion
    def ver_marca(codigo: str):
        conexion = conexion_actual()
        ficha = ficha_marca(conexion, codigo)
        if not ficha:
            return _registro_no_encontrado("No se encontró esa marca.")
        ficha["marca"]["imagen_url"] = url_imagen(ficha["marca"]["id"], ficha["marca"].get("archivo_png"))
        for acompanante in ficha["acompanantes"]:
            acompanante["imagen_url"] = url_imagen(acompanante["id"], acompanante.get("archivo_png"))
        nombres = nombres_usuarios(
            conexion, [ficha["marca"].get("creado_por"), ficha["marca"].get("actualizado_por")]
        )
        ficha["marca"]["creado_por_nombre"] = nombres.get(str(ficha["marca"].get("creado_por")))
        ficha["marca"]["actualizado_por_nombre"] = nombres.get(str(ficha["marca"].get("actualizado_por")))
        cambio_pendiente = cambio_pendiente_de(conexion, "marcas", ficha["marca"]["id"])
        return render_template(
            "marca_detalle.html",
            activo="buscar",
            ficha=ficha,
            cambio_pendiente=cambio_pendiente,
            perfil=perfil_actual(),
        )

    @app.post("/marcas/<int:marca_id>")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def editar_marca(marca_id: int):
        conexion = conexion_actual()
        perfil = perfil_actual()
        marca_actual = obtener_marca_por_id(conexion, marca_id)
        if not marca_actual:
            return _registro_no_encontrado("No se encontró esa marca.")

        cambios: dict = {}
        for campo in CAMPOS_MARCA_EDITABLES:
            if campo in ("archivo_png", "archivo_svg"):
                continue
            if campo in request.form:
                valor = request.form.get(campo, "").strip() or None
                if campo == "codigo":
                    if not valor:
                        continue  # el código no puede quedar vacío
                    if not _CODIGO_VALIDO.match(valor):
                        return (
                            "El código sólo puede tener letras, números, puntos, "
                            "guiones y guiones bajos (se usa también como parte de la "
                            "dirección web y del nombre de archivo).",
                            400,
                        )
                if valor != marca_actual.get(campo):
                    cambios[campo] = valor

        # Las imágenes nuevas se leen acá pero recién se aplican más abajo,
        # según el rol: un Operador las escribe directo; para un
        # Administrador quedan "en espera" (codificadas en el propio cambio
        # propuesto) hasta que otra persona lo apruebe -- así una propuesta
        # rechazada nunca llega a pisar la imagen real.
        archivos_subidos: dict[str, bytes] = {}
        for campo, extension in (("nuevo_png", "png"), ("nuevo_svg", "svg")):
            archivo = request.files.get(campo)
            if archivo and archivo.filename:
                archivos_subidos[extension] = archivo.read()

        codigo_para_volver = marca_actual["codigo"]
        if cambios or archivos_subidos:
            try:
                if perfil["rol"] == "operador":
                    for extension, contenido in archivos_subidos.items():
                        subir_imagen_marca(conexion, marca_id, contenido, extension)
                    if cambios:
                        actualizar_marca_campos(conexion, marca_id, cambios)
                    codigo_para_volver = cambios.get("codigo", codigo_para_volver)
                else:  # administrador -- pasa por modificación supervisada
                    valores_anteriores = {k: marca_actual.get(k) for k in cambios}
                    for extension, contenido in archivos_subidos.items():
                        campo = f"archivo_{extension}"
                        cambios[campo] = base64.b64encode(contenido).decode("ascii")
                        valores_anteriores[campo] = None  # no se arrastra la imagen vieja completa a la cola
                    proponer_cambio(conexion, "marcas", marca_id, cambios, valores_anteriores, perfil["id"])
            except Exception as exc:
                return _error_seguro("No se pudo guardar el cambio.", exc)
        return redirect(url_for("ver_marca", codigo=codigo_para_volver))

    @app.post("/marcas/<int:marca_id>/ocultar")
    @requiere_sesion
    @requiere_rol("administrador")
    def ocultar_marca_ruta(marca_id: int):
        conexion = conexion_actual()
        marca = obtener_marca_por_id(conexion, marca_id)
        if not marca:
            return _registro_no_encontrado("No se encontró esa marca.")
        ocultar_marca(conexion, marca_id)
        return redirect(url_for("ver_marca", codigo=marca["codigo"]))

    @app.post("/marcas/<int:marca_id>/restaurar")
    @requiere_sesion
    @requiere_rol("administrador")
    def restaurar_marca_ruta(marca_id: int):
        conexion = conexion_actual()
        marca = obtener_marca_por_id(conexion, marca_id)
        if not marca:
            return _registro_no_encontrado("No se encontró esa marca.")
        restaurar_marca(conexion, marca_id)
        return redirect(url_for("ver_marca", codigo=marca["codigo"]))

    @app.get("/imagenes/marca/<int:marca_id>.<extension>", endpoint="imagen_marca")
    @requiere_sesion
    def imagen_marca_ruta(marca_id: int, extension: str):
        if extension not in ("png", "svg"):
            return _registro_no_encontrado("Formato de imagen no soportado.")
        conexion = conexion_actual()
        contenido = imagen_marca(conexion, marca_id, extension)
        if not contenido:
            return _registro_no_encontrado("Esa marca no tiene esa imagen guardada.")
        mimetype = "image/svg+xml" if extension == "svg" else "image/png"
        return send_file(io.BytesIO(contenido), mimetype=mimetype, max_age=3600)

    @app.get("/guias")
    @requiere_sesion
    def guias():
        conexion = conexion_actual()
        perfil = perfil_actual()
        pagina = max(1, request.args.get("pagina", 1, type=int))
        texto = request.args.get("q") or None
        estado = request.args.get("estado") or None
        creada_desde = request.args.get("desde") or None
        creada_hasta = request.args.get("hasta") or None
        ver_ocultos = perfil["rol"] == "administrador" and request.args.get("ver") == "ocultos"
        operaciones, total = listar_operaciones_paginado(
            conexion, pagina=pagina, texto=texto, estado=estado,
            creada_desde=creada_desde, creada_hasta=creada_hasta, tipo_operacion="compra",
            incluir_ocultos=ver_ocultos,
        )
        return render_template(
            "guias.html",
            activo="guias",
            modo="compra",
            operaciones=operaciones,
            pagina=pagina,
            total=total,
            texto=texto or "",
            estado=estado or "",
            desde=creada_desde or "",
            hasta=creada_hasta or "",
            ver_ocultos=ver_ocultos,
            perfil=perfil,
            filtros_guardados=listar_filtros(conexion, perfil["id"], "guias"),
            columnas_visibles=preferencias_columnas(conexion, perfil["id"], "guias"),
        )

    @app.get("/ventas")
    @requiere_sesion
    def ventas():
        conexion = conexion_actual()
        perfil = perfil_actual()
        pagina = max(1, request.args.get("pagina", 1, type=int))
        texto = request.args.get("q") or None
        estado = request.args.get("estado") or None
        creada_desde = request.args.get("desde") or None
        creada_hasta = request.args.get("hasta") or None
        ver_ocultos = perfil["rol"] == "administrador" and request.args.get("ver") == "ocultos"
        operaciones, total = listar_operaciones_paginado(
            conexion, pagina=pagina, texto=texto, estado=estado,
            creada_desde=creada_desde, creada_hasta=creada_hasta, tipo_operacion="venta",
            incluir_ocultos=ver_ocultos,
        )
        return render_template(
            "guias.html",
            activo="ventas",
            modo="venta",
            operaciones=operaciones,
            pagina=pagina,
            total=total,
            texto=texto or "",
            estado=estado or "",
            desde=creada_desde or "",
            hasta=creada_hasta or "",
            ver_ocultos=ver_ocultos,
            perfil=perfil,
            filtros_guardados=listar_filtros(conexion, perfil["id"], "ventas"),
            columnas_visibles=preferencias_columnas(conexion, perfil["id"], "ventas"),
        )

    @app.get("/exportar")
    @requiere_sesion
    def exportar():
        conexion = conexion_actual()
        seccion = request.args.get("seccion") or "operaciones"
        if seccion not in ("operaciones", "marcas"):
            seccion = "operaciones"
        texto = request.args.get("q") or None
        estado = request.args.get("estado") or None
        resultados = []
        total = 0
        if seccion == "marcas":
            tipo = request.args.get("tipo") or None
            resultados, total = buscar_marcas(
                conexion, texto, pagina=1, por_pagina=LIMITE_EXPORTAR, estado=estado, tipo=tipo
            )
            filtros = {"q": texto or "", "estado": estado or "", "tipo": tipo or ""}
        else:
            tipo_operacion = request.args.get("tipo_operacion") or None
            if tipo_operacion not in ("compra", "venta"):
                tipo_operacion = None
            desde = request.args.get("desde") or None
            hasta = request.args.get("hasta") or None
            resultados, total = listar_operaciones_paginado(
                conexion, pagina=1, por_pagina=LIMITE_EXPORTAR, texto=texto, estado=estado,
                creada_desde=desde, creada_hasta=hasta, tipo_operacion=tipo_operacion,
            )
            filtros = {
                "q": texto or "", "estado": estado or "", "tipo_operacion": tipo_operacion or "",
                "desde": desde or "", "hasta": hasta or "",
            }
        return render_template(
            "exportar.html",
            activo="exportar",
            seccion=seccion,
            resultados=resultados,
            total=total,
            limite=LIMITE_EXPORTAR,
            filtros=filtros,
            perfil=perfil_actual(),
        )

    @app.post("/exportar/operaciones.xlsx")
    @requiere_sesion
    def exportar_operaciones_xlsx():
        conexion = conexion_actual()
        ids = [int(v) for v in request.form.getlist("ids") if v.isdigit()]
        if not ids:
            return _error_seguro("Marcá al menos un registro para exportar.", ValueError("sin ids"))
        try:
            filas = exportar_operaciones(conexion, ids=ids)
        except Exception as exc:
            return _error_seguro("No se pudo generar la planilla.", exc)
        etiquetas_tipo = {"compra": "Compra", "venta": "Venta"}
        cuerpo = [
            [
                etiquetas_tipo.get(o.get("tipo_operacion"), o.get("tipo_operacion") or ""),
                o.get("numero_guia") or "", o.get("fecha") or "", o.get("vendedor_nombre") or "",
                o.get("vendedor_documento") or "", o.get("comprador_nombre") or "",
                o.get("comprador_documento") or "",
                o.get("cantidad_animales") if o.get("cantidad_animales") is not None else "",
                o.get("categoria_animales") or "", o.get("revisar") or "",
                "Sí" if o.get("guia_colisionada") else "No", _fecha_sin_tz(o.get("creado_en")),
            ]
            for o in filas
        ]
        buffer = _libro_excel(
            "Guías y ventas",
            ["Tipo", "N° de guía", "Fecha", "Vendedor", "Doc. vendedor", "Comprador", "Doc. comprador",
             "Cantidad", "Categoría", "Revisar", "Colisiona", "Creado en"],
            cuerpo,
        )
        return send_file(
            buffer, as_attachment=True, download_name="guias_y_ventas.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.post("/exportar/marcas.xlsx")
    @requiere_sesion
    def exportar_marcas_xlsx():
        conexion = conexion_actual()
        ids = [int(v) for v in request.form.getlist("ids") if v.isdigit()]
        if not ids:
            return _error_seguro("Marcá al menos un registro para exportar.", ValueError("sin ids"))
        try:
            filas = exportar_marcas(conexion, ids=ids)
        except Exception as exc:
            return _error_seguro("No se pudo generar la planilla.", exc)
        etiquetas_tipo = {"dominante": "Dominante", "complementaria": "Complementaria"}
        etiquetas_estado = {"activa": "Al día", "revisar": "Revisar", "baja": "De baja"}
        cuerpo = []
        for m in filas:
            propietario = m.get("propietarios") or {}
            operacion = m.get("operaciones") or {}
            cuerpo.append([
                m.get("codigo") or "", etiquetas_tipo.get(m.get("tipo"), m.get("tipo") or ""),
                etiquetas_estado.get(m.get("estado"), m.get("estado") or ""), m.get("vence_en") or "",
                propietario.get("nombre") or "", propietario.get("documento") or "",
                operacion.get("numero_guia") or "", operacion.get("fecha") or "",
            ])
        buffer = _libro_excel(
            "Marcas",
            ["Código", "Tipo", "Estado", "Vence", "Propietario", "Documento", "N° de guía", "Fecha"],
            cuerpo,
        )
        return send_file(
            buffer, as_attachment=True, download_name="marcas.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    _CAMPOS_FILTRO_POR_SECCION = {
        "marcas": ("q", "estado", "tipo"),
        "guias": ("q", "estado", "desde", "hasta"),
        "ventas": ("q", "estado", "desde", "hasta"),
    }

    @app.post("/filtros/<seccion>")
    @requiere_sesion
    def guardar_filtro_ruta(seccion: str):
        if seccion not in SECCIONES_FILTRO:
            return _registro_no_encontrado("Sección de filtros no válida.")
        volver = request.form.get("volver") or url_for("buscar")
        nombre = (request.form.get("nombre") or "").strip()
        if not nombre:
            return redirect(volver)
        conexion = conexion_actual()
        perfil = perfil_actual()
        campos = _CAMPOS_FILTRO_POR_SECCION[seccion]
        parametros = {c: request.form.get(c) for c in campos if request.form.get(c)}
        guardar_filtro(conexion, perfil["id"], seccion, nombre, parametros)
        return redirect(volver)

    @app.post("/filtros/<int:filtro_id>/eliminar")
    @requiere_sesion
    def eliminar_filtro_ruta(filtro_id: int):
        conexion = conexion_actual()
        perfil = perfil_actual()
        eliminar_filtro(conexion, perfil["id"], filtro_id)
        volver = request.form.get("volver") or url_for("buscar")
        return redirect(volver)

    @app.post("/preferencias/columnas/<seccion>")
    @requiere_sesion
    def actualizar_preferencias_columnas_ruta(seccion: str):
        if seccion not in SECCIONES_FILTRO:
            return _registro_no_encontrado("Sección no válida.")
        conexion = conexion_actual()
        perfil = perfil_actual()
        columnas = request.form.getlist("columna")
        actualizar_preferencias_columnas(conexion, perfil["id"], seccion, columnas)
        volver = request.form.get("volver") or url_for("buscar")
        return redirect(volver)

    @app.get("/guias/nueva")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def nueva_guia():
        return render_template("guia_nueva.html", activo="guias", modo="compra", perfil=perfil_actual())

    @app.get("/ventas/nueva")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def nueva_venta():
        return render_template("guia_nueva.html", activo="ventas", modo="venta", perfil=perfil_actual())

    @app.post("/guias/nueva")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def crear_guia():
        conexion = conexion_actual()
        perfil = perfil_actual()
        campos = {}
        for campo in CAMPOS_OPERACION_EDITABLES:
            valor = request.form.get(campo, "").strip() or None
            if campo == "cantidad_animales" and valor is not None:
                try:
                    valor = int(valor)
                except ValueError:
                    valor = None
            campos[campo] = valor
        campos["tipo_operacion"] = "compra"
        try:
            operacion_id = crear_operacion(conexion, campos, perfil["id"])
        except Exception as exc:
            return _error_seguro("No se pudo crear la guía.", exc)
        return redirect(url_for("ver_guia", operacion_id=operacion_id))

    @app.post("/ventas/nueva")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def crear_venta():
        conexion = conexion_actual()
        perfil = perfil_actual()
        campos = {}
        for campo in CAMPOS_OPERACION_EDITABLES:
            valor = request.form.get(campo, "").strip() or None
            if campo == "cantidad_animales" and valor is not None:
                try:
                    valor = int(valor)
                except ValueError:
                    valor = None
            campos[campo] = valor
        campos["tipo_operacion"] = "venta"
        try:
            operacion_id = crear_operacion(conexion, campos, perfil["id"])
        except Exception as exc:
            return _error_seguro("No se pudo crear la venta.", exc)
        return redirect(url_for("ver_guia", operacion_id=operacion_id))

    @app.get("/guias/<int:operacion_id>")
    @requiere_sesion
    def ver_guia(operacion_id: int):
        conexion = conexion_actual()
        operacion = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion:
            return _registro_no_encontrado("No se encontró esa guía.")
        marcas = marcas_de_operacion(conexion, operacion_id)
        for m in marcas:
            m["imagen_url"] = url_imagen(m["id"], m.get("archivo_png"))
        nombres = nombres_usuarios(conexion, [operacion.get("creado_por")])
        operacion["creado_por_nombre"] = nombres.get(str(operacion.get("creado_por")))
        cambio_pendiente = cambio_pendiente_de(conexion, "operaciones", operacion_id)
        return render_template(
            "guia_detalle.html",
            activo="ventas" if operacion.get("tipo_operacion") == "venta" else "guias",
            operacion=operacion,
            marcas=marcas,
            cambio_pendiente=cambio_pendiente,
            perfil=perfil_actual(),
        )

    @app.post("/guias/<int:operacion_id>")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def editar_guia(operacion_id: int):
        conexion = conexion_actual()
        perfil = perfil_actual()
        operacion_actual = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion_actual:
            return _registro_no_encontrado("No se encontró esa guía.")

        cambios: dict = {}
        for campo in CAMPOS_OPERACION_EDITABLES:
            if campo not in request.form:
                continue
            valor = request.form.get(campo, "").strip() or None
            if campo == "cantidad_animales" and valor is not None:
                try:
                    valor = int(valor)
                except ValueError:
                    continue
            if valor != operacion_actual.get(campo):
                cambios[campo] = valor

        if cambios:
            try:
                if perfil["rol"] == "operador":
                    actualizar_operacion_campos(conexion, operacion_id, cambios)
                else:  # administrador -- pasa por modificación supervisada
                    valores_anteriores = {k: operacion_actual.get(k) for k in cambios}
                    proponer_cambio(conexion, "operaciones", operacion_id, cambios, valores_anteriores, perfil["id"])
            except Exception as exc:
                return _error_seguro("No se pudo guardar el cambio.", exc)
        return redirect(url_for("ver_guia", operacion_id=operacion_id))

    @app.post("/guias/<int:operacion_id>/ocultar")
    @requiere_sesion
    @requiere_rol("administrador")
    def ocultar_guia_ruta(operacion_id: int):
        conexion = conexion_actual()
        if not obtener_operacion_por_id(conexion, operacion_id):
            return _registro_no_encontrado("No se encontró esa guía.")
        ocultar_operacion(conexion, operacion_id)
        return redirect(url_for("ver_guia", operacion_id=operacion_id))

    @app.post("/guias/<int:operacion_id>/restaurar")
    @requiere_sesion
    @requiere_rol("administrador")
    def restaurar_guia_ruta(operacion_id: int):
        conexion = conexion_actual()
        if not obtener_operacion_por_id(conexion, operacion_id):
            return _registro_no_encontrado("No se encontró esa guía.")
        restaurar_operacion(conexion, operacion_id)
        return redirect(url_for("ver_guia", operacion_id=operacion_id))

    @app.post("/guias/<int:operacion_id>/marcas")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def agregar_marca_a_guia(operacion_id: int):
        conexion = conexion_actual()
        perfil = perfil_actual()
        operacion = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion:
            return _registro_no_encontrado("No se encontró esa guía.")

        tipo = request.form.get("tipo") or "complementaria"
        descripcion = request.form.get("descripcion", "").strip() or None
        # Con "multiple" en el campo de archivo se puede cargar de una vez
        # todo un grupo de marcas complementarias (una por imagen); sin
        # ninguna imagen se agrega una sola marca en blanco, como antes.
        archivos = [a for a in request.files.getlist("imagen") if a and a.filename]

        try:
            for archivo in archivos or [None]:
                nueva = crear_marca(
                    conexion, {"tipo": tipo, "descripcion": descripcion, "operacion_id": operacion_id}, perfil["id"]
                )
                if archivo:
                    extension = "svg" if archivo.filename.lower().endswith(".svg") else "png"
                    subir_imagen_marca(conexion, nueva["id"], archivo.read(), extension)
        except Exception as exc:
            return _error_seguro("No se pudo agregar la marca.", exc)

        return redirect(url_for("ver_guia", operacion_id=operacion_id))

    @app.get("/guias/<int:operacion_id>/imprimir")
    @requiere_sesion
    def imprimir_guia(operacion_id: int):
        conexion = conexion_actual()
        operacion = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion:
            return _registro_no_encontrado("No se encontró esa guía.")
        marcas = marcas_de_operacion(conexion, operacion_id)
        for m in marcas:
            m["imagen_url"] = url_imagen(m["id"], m.get("archivo_png"))
        borrador = operacion.get("borrador_venta")
        ids_elegidas_borrador = {int(v) for v in (borrador or {}).get("marcas", []) if v.isdigit()}
        return render_template(
            "guia_imprimir.html",
            activo="guias",
            operacion=operacion,
            marcas=marcas,
            tiene_borrador_pdf=bool(operacion.get("borrador_pdf")),
            borrador_pdf_nombre=operacion.get("borrador_pdf_nombre"),
            ids_elegidas_borrador=ids_elegidas_borrador,
            perfil=perfil_actual(),
        )

    @app.post("/guias/<int:operacion_id>/imprimir")
    @requiere_sesion
    def generar_guia_pdf(operacion_id: int):
        conexion = conexion_actual()
        operacion = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion:
            return _registro_no_encontrado("No se encontró esa guía.")

        archivo_pdf = request.files.get("pdf_guia")
        if not archivo_pdf or not archivo_pdf.filename:
            return "Subí el PDF de la guía descargado de SENACSA.", 400

        ids_elegidos = {int(v) for v in request.form.getlist("marca_id")}
        if not ids_elegidos:
            return "Seleccioná al menos una marca.", 400

        marcas = [m for m in marcas_de_operacion(conexion, operacion_id) if m["id"] in ids_elegidos]

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            pdf_entrada = tmp_path / "entrada.pdf"
            archivo_pdf.save(pdf_entrada)

            rutas_imagenes = []
            for m in marcas:
                contenido = imagen_marca(conexion, m["id"], "png")
                if not contenido:
                    continue
                destino = tmp_path / f"{m['codigo']}.png"
                destino.write_bytes(contenido)
                rutas_imagenes.append(destino)

            if not rutas_imagenes:
                return "Ninguna de las marcas seleccionadas tiene una imagen disponible.", 400

            salida = tmp_path / "salida.pdf"
            try:
                completar_guia(pdf_entrada, rutas_imagenes, salida)
            except Exception as exc:
                return _error_seguro("No se pudo generar el PDF.", exc)
            contenido_pdf = salida.read_bytes()

        if operacion.get("borrador_venta"):
            try:
                eliminar_borrador_venta(conexion, operacion_id)
            except Exception as exc:
                print(f"generar_guia_pdf: no se pudo limpiar el borrador de {operacion_id}: {exc}")

        nombre_descarga = f"guia_{operacion.get('numero_guia') or operacion_id}.pdf"
        return send_file(
            io.BytesIO(contenido_pdf),
            as_attachment=True,
            download_name=nombre_descarga,
            mimetype="application/pdf",
        )

    @app.post("/guias/<int:operacion_id>/borrador")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def guardar_borrador_guia_ruta(operacion_id: int):
        """Mismo mecanismo que el borrador de Venta (guardar_borrador_venta
        es genérica por operacion_id) -- acá la lista guardada son ids de
        marcas ya cargadas en esta guía, no códigos del catálogo entero."""
        conexion = conexion_actual()
        if not obtener_operacion_por_id(conexion, operacion_id):
            return _registro_no_encontrado("No se encontró esa guía.")
        ids = [v.strip() for v in request.form.getlist("marca_id") if v.strip()]
        archivo_pdf = request.files.get("pdf_guia")
        pdf_bytes = archivo_pdf.read() if archivo_pdf and archivo_pdf.filename else None
        pdf_nombre = archivo_pdf.filename if archivo_pdf and archivo_pdf.filename else None
        guardar_borrador_venta(
            conexion, operacion_id, ids, datetime.now(timezone.utc).isoformat(),
            pdf_bytes=pdf_bytes, pdf_nombre=pdf_nombre,
        )
        return redirect(url_for("imprimir_guia", operacion_id=operacion_id))

    @app.post("/guias/<int:operacion_id>/borrador/descartar")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def descartar_borrador_guia_ruta(operacion_id: int):
        conexion = conexion_actual()
        if not obtener_operacion_por_id(conexion, operacion_id):
            return _registro_no_encontrado("No se encontró esa guía.")
        eliminar_borrador_venta(conexion, operacion_id)
        return redirect(url_for("imprimir_guia", operacion_id=operacion_id))

    @app.get("/guias/<int:operacion_id>/borrador/pdf")
    @requiere_sesion
    def borrador_guia_pdf(operacion_id: int):
        conexion = conexion_actual()
        resultado = obtener_borrador_venta_pdf(conexion, operacion_id)
        if not resultado:
            return _registro_no_encontrado("Ese borrador no tiene un PDF guardado.")
        contenido, nombre = resultado
        return send_file(io.BytesIO(contenido), download_name=nombre, mimetype="application/pdf")

    @app.post("/guias/<int:operacion_id>/previsualizar")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def previsualizar_guia(operacion_id: int):
        conexion = conexion_actual()
        operacion = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion:
            return jsonify({"error": "No se encontró esa guía."}), 404

        archivo_pdf = request.files.get("pdf_guia")
        pdf_bytes = archivo_pdf.read() if archivo_pdf and archivo_pdf.filename else None
        pdf_nombre = archivo_pdf.filename if archivo_pdf and archivo_pdf.filename else None

        ids = [v.strip() for v in request.form.getlist("marca_id") if v.strip()]
        if pdf_bytes or ids != (operacion.get("borrador_venta") or {}).get("marcas", []):
            guardar_borrador_venta(
                conexion, operacion_id, ids, datetime.now(timezone.utc).isoformat(),
                pdf_bytes=pdf_bytes, pdf_nombre=pdf_nombre,
            )

        if not pdf_bytes:
            guardado = obtener_borrador_venta_pdf(conexion, operacion_id)
            if not guardado:
                return jsonify({"error": "Subí primero el PDF de la guía."}), 400
            pdf_bytes = guardado[0]

        ids_int = {int(v) for v in ids if v.isdigit()}
        marcas_elegidas = [m for m in marcas_de_operacion(conexion, operacion_id) if m["id"] in ids_int]

        try:
            resultado = _previsualizar_estampado(conexion, operacion, marcas_elegidas, pdf_bytes)
        except Exception as exc:
            print(f"previsualizar_guia: {exc}")
            return jsonify({"error": "No se pudo generar la vista previa."}), 400
        return jsonify(resultado)

    @app.get("/marcas/buscar.json")
    @requiere_sesion
    def buscar_marcas_json():
        """Para el buscador de marcas complementarias al armar una Venta --
        busca en TODO el catálogo existente (no en una guía puntual), porque
        lo que se vende ya está cargado de antes."""
        conexion = conexion_actual()
        texto = request.args.get("q") or None
        resultados, total = buscar_marcas(conexion, texto, pagina=1, por_pagina=15)
        return jsonify([
            {
                "codigo": m["codigo"],
                "tipo": m.get("tipo"),
                "estado": m.get("estado"),
                "propietario": (m.get("propietarios") or {}).get("nombre"),
                "numero_guia": (m.get("operaciones") or {}).get("numero_guia"),
                "imagen_url": url_imagen(m["id"], m.get("archivo_png")),
            }
            for m in resultados
        ])

    @app.get("/ventas/<int:operacion_id>/imprimir")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def imprimir_venta(operacion_id: int):
        conexion = conexion_actual()
        operacion = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion or operacion.get("tipo_operacion") != "venta":
            return _registro_no_encontrado("No se encontró esa venta.")

        borrador = operacion.get("borrador_venta")
        marcas_borrador = []
        if borrador and borrador.get("marcas"):
            por_codigo = marcas_por_codigos(conexion, borrador["marcas"])
            marcas_borrador = [
                {
                    "codigo": codigo,
                    "tipo": por_codigo[codigo].get("tipo"),
                    "propietario": (por_codigo[codigo].get("propietarios") or {}).get("nombre"),
                    "imagen_url": url_imagen(por_codigo[codigo]["id"], por_codigo[codigo].get("archivo_png")),
                }
                for codigo in borrador["marcas"]
                if codigo in por_codigo
            ]
        borrador_para_plantilla = None
        if borrador:
            borrador_para_plantilla = {
                **borrador,
                "pdf_ruta": bool(operacion.get("borrador_pdf")),
                "pdf_nombre_original": operacion.get("borrador_pdf_nombre"),
                "dominante_ruta": bool(operacion.get("borrador_dominante_png")),
            }
        return render_template(
            "venta_imprimir.html",
            activo="ventas",
            operacion=operacion,
            perfil=perfil_actual(),
            borrador=borrador_para_plantilla,
            marcas_borrador=marcas_borrador,
        )

    @app.post("/ventas/<int:operacion_id>/borrador")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def guardar_borrador_venta_ruta(operacion_id: int):
        """Guarda el avance de la carga (marcas elegidas, PDF y Dominante ya
        puestos) para que la persona pueda cerrar esto y continuar más tarde
        -- desde cualquier computadora, no sólo desde este navegador."""
        conexion = conexion_actual()
        operacion = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion or operacion.get("tipo_operacion") != "venta":
            return _registro_no_encontrado("No se encontró esa venta.")

        marcas = [c.strip() for c in request.form.getlist("marca_codigo") if c.strip()]
        archivo_pdf = request.files.get("pdf_guia")
        pdf_bytes = archivo_pdf.read() if archivo_pdf and archivo_pdf.filename else None
        pdf_nombre = archivo_pdf.filename if archivo_pdf and archivo_pdf.filename else None
        archivo_dominante = request.files.get("imagen_dominante")
        dominante_bytes = archivo_dominante.read() if archivo_dominante and archivo_dominante.filename else None

        guardar_borrador_venta(
            conexion, operacion_id, marcas, datetime.now(timezone.utc).isoformat(),
            pdf_bytes=pdf_bytes, pdf_nombre=pdf_nombre, dominante_bytes=dominante_bytes,
        )
        return redirect(url_for("imprimir_venta", operacion_id=operacion_id))

    @app.post("/ventas/<int:operacion_id>/borrador/descartar")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def descartar_borrador_venta_ruta(operacion_id: int):
        conexion = conexion_actual()
        operacion = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion or operacion.get("tipo_operacion") != "venta":
            return _registro_no_encontrado("No se encontró esa venta.")
        eliminar_borrador_venta(conexion, operacion_id)
        return redirect(url_for("imprimir_venta", operacion_id=operacion_id))

    @app.get("/ventas/<int:operacion_id>/borrador/pdf")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def borrador_venta_pdf(operacion_id: int):
        conexion = conexion_actual()
        resultado = obtener_borrador_venta_pdf(conexion, operacion_id)
        if not resultado:
            return _registro_no_encontrado("Ese borrador no tiene un PDF guardado.")
        contenido, nombre = resultado
        return send_file(io.BytesIO(contenido), download_name=nombre, mimetype="application/pdf")

    @app.get("/ventas/<int:operacion_id>/borrador/dominante.png")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def borrador_venta_dominante(operacion_id: int):
        conexion = conexion_actual()
        contenido = obtener_borrador_venta_dominante(conexion, operacion_id)
        if not contenido:
            return _registro_no_encontrado("Ese borrador no tiene una Dominante guardada.")
        return send_file(io.BytesIO(contenido), download_name="dominante.png", mimetype="image/png")

    @app.post("/ventas/<int:operacion_id>/imprimir")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def generar_venta_pdf(operacion_id: int):
        conexion = conexion_actual()
        operacion = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion or operacion.get("tipo_operacion") != "venta":
            return _registro_no_encontrado("No se encontró esa venta.")

        archivo_pdf = request.files.get("pdf_guia")
        if not archivo_pdf or not archivo_pdf.filename:
            return _error_seguro("Subí el PDF de la guía descargado de SENACSA.", "sin archivo")

        codigos_elegidos = [c.strip() for c in request.form.getlist("marca_codigo") if c.strip()]
        if not codigos_elegidos:
            return _error_seguro("Elegí al menos una marca complementaria para la venta.", "sin marcas")

        por_codigo = marcas_por_codigos(conexion, codigos_elegidos)
        faltantes = [c for c in codigos_elegidos if c not in por_codigo]
        if faltantes:
            return _error_seguro(f"No se encontró la marca {faltantes[0]} en el catálogo.", "código inexistente")

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            pdf_entrada = tmp_path / "entrada.pdf"
            archivo_pdf.save(pdf_entrada)

            imagen_dominante = None
            archivo_dominante = request.files.get("imagen_dominante")
            if archivo_dominante and archivo_dominante.filename:
                imagen_dominante = tmp_path / "dominante.png"
                archivo_dominante.save(imagen_dominante)

            rutas_complementarias = []
            for codigo in codigos_elegidos:
                marca = por_codigo[codigo]
                contenido = imagen_marca(conexion, marca["id"], "png")
                if not contenido:
                    continue
                destino = tmp_path / f"{codigo}.png"
                destino.write_bytes(contenido)
                rutas_complementarias.append(destino)

            if not rutas_complementarias:
                return _error_seguro("Ninguna de las marcas elegidas tiene una imagen PNG disponible.", "sin imágenes")

            salida = tmp_path / "salida.pdf"
            try:
                completar_venta(pdf_entrada, imagen_dominante, rutas_complementarias, salida)
            except Exception as exc:
                return _error_seguro("No se pudo generar el PDF de la venta.", exc)
            contenido_pdf = salida.read_bytes()

        if operacion.get("borrador_venta"):
            try:
                eliminar_borrador_venta(conexion, operacion_id)
            except Exception as exc:
                print(f"generar_venta_pdf: no se pudo limpiar el borrador de {operacion_id}: {exc}")

        nombre_descarga = f"venta_{operacion.get('numero_guia') or operacion_id}.pdf"
        return send_file(
            io.BytesIO(contenido_pdf),
            as_attachment=True,
            download_name=nombre_descarga,
            mimetype="application/pdf",
        )

    @app.post("/ventas/<int:operacion_id>/previsualizar")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def previsualizar_venta(operacion_id: int):
        """Vista previa en vivo del panel izquierdo: corre completar_venta
        sobre un archivo temporal con las marcas elegidas hasta el momento
        (nunca sobre el PDF final) y devuelve las páginas como imágenes,
        más el cupo libre y los datos leídos del encabezado. El PDF y la
        Dominante recién subidos quedan guardados como borrador para no
        tener que volver a mandarlos en cada llamada."""
        conexion = conexion_actual()
        operacion = obtener_operacion_por_id(conexion, operacion_id)
        if not operacion or operacion.get("tipo_operacion") != "venta":
            return jsonify({"error": "No se encontró esa venta."}), 404

        archivo_pdf = request.files.get("pdf_guia")
        pdf_bytes = archivo_pdf.read() if archivo_pdf and archivo_pdf.filename else None
        pdf_nombre = archivo_pdf.filename if archivo_pdf and archivo_pdf.filename else None
        archivo_dominante = request.files.get("imagen_dominante")
        dominante_bytes = archivo_dominante.read() if archivo_dominante and archivo_dominante.filename else None

        codigos = [c.strip() for c in request.form.getlist("marca_codigo") if c.strip()]
        if pdf_bytes or dominante_bytes or codigos != (operacion.get("borrador_venta") or {}).get("marcas", []):
            guardar_borrador_venta(
                conexion, operacion_id, codigos, datetime.now(timezone.utc).isoformat(),
                pdf_bytes=pdf_bytes, pdf_nombre=pdf_nombre, dominante_bytes=dominante_bytes,
            )
            operacion = obtener_operacion_por_id(conexion, operacion_id)

        if not pdf_bytes:
            guardado = obtener_borrador_venta_pdf(conexion, operacion_id)
            if not guardado:
                return jsonify({"error": "Subí primero el PDF de la guía."}), 400
            pdf_bytes = guardado[0]
        if not dominante_bytes:
            dominante_bytes = obtener_borrador_venta_dominante(conexion, operacion_id)

        por_codigo = marcas_por_codigos(conexion, codigos)
        marcas_elegidas = [por_codigo[c] for c in codigos if c in por_codigo]

        try:
            resultado = _previsualizar_estampado(conexion, operacion, marcas_elegidas, pdf_bytes, dominante_bytes)
        except Exception as exc:
            print(f"previsualizar_venta: {exc}")
            return jsonify({"error": "No se pudo generar la vista previa."}), 400
        return jsonify(resultado)

    @app.get("/propietarios")
    @requiere_sesion
    def propietarios():
        conexion = conexion_actual()
        perfil = perfil_actual()
        texto = request.args.get("q") or None
        pagina = max(1, request.args.get("pagina", 1, type=int))
        ver_ocultos = perfil["rol"] == "administrador" and request.args.get("ver") == "ocultos"
        resultados, total = listar_propietarios(
            conexion, texto, pagina=pagina, incluir_ocultos=ver_ocultos
        )
        return render_template(
            "propietarios.html",
            activo="propietarios",
            resultados=resultados,
            total=total,
            pagina=pagina,
            por_pagina=POR_PAGINA_PROPIETARIOS,
            texto=texto or "",
            ver_ocultos=ver_ocultos,
            perfil=perfil,
        )

    @app.get("/propietarios/<int:propietario_id>")
    @requiere_sesion
    def ver_propietario(propietario_id: int):
        conexion = conexion_actual()
        propietario = obtener_propietario(conexion, propietario_id)
        if not propietario:
            return _registro_no_encontrado("No se encontró ese propietario.")
        marcas = marcas_de_propietario(conexion, propietario_id)
        for m in marcas:
            m["imagen_url"] = url_imagen(m["id"], m.get("archivo_png"))
        operaciones = operaciones_de_propietario(conexion, propietario.get("documento"))
        return render_template(
            "propietario_detalle.html",
            activo="propietarios",
            propietario=propietario,
            marcas=marcas,
            operaciones=operaciones,
            perfil=perfil_actual(),
        )

    @app.post("/propietarios/<int:propietario_id>")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def editar_propietario(propietario_id: int):
        conexion = conexion_actual()
        propietario_actual = obtener_propietario(conexion, propietario_id)
        if not propietario_actual:
            return _registro_no_encontrado("No se encontró ese propietario.")
        campos = {}
        for campo in CAMPOS_PROPIETARIO_EDITABLES:
            if campo not in request.form:
                continue
            valor = request.form.get(campo, "").strip() or None
            if valor != propietario_actual.get(campo):
                campos[campo] = valor
        if campos:
            try:
                actualizar_propietario(conexion, propietario_id, campos)
            except Exception as exc:
                return _error_seguro("No se pudo guardar el cambio.", exc)
        return redirect(url_for("ver_propietario", propietario_id=propietario_id))

    @app.post("/propietarios/<int:propietario_id>/ocultar")
    @requiere_sesion
    @requiere_rol("administrador")
    def ocultar_propietario_ruta(propietario_id: int):
        conexion = conexion_actual()
        if not obtener_propietario(conexion, propietario_id):
            return _registro_no_encontrado("No se encontró ese propietario.")
        ocultar_propietario(conexion, propietario_id)
        return redirect(url_for("ver_propietario", propietario_id=propietario_id))

    @app.post("/propietarios/<int:propietario_id>/restaurar")
    @requiere_sesion
    @requiere_rol("administrador")
    def restaurar_propietario_ruta(propietario_id: int):
        conexion = conexion_actual()
        if not obtener_propietario(conexion, propietario_id):
            return _registro_no_encontrado("No se encontró ese propietario.")
        restaurar_propietario(conexion, propietario_id)
        return redirect(url_for("ver_propietario", propietario_id=propietario_id))

    @app.get("/estadisticas")
    @requiere_sesion
    def estadisticas_vista():
        conexion = conexion_actual()
        top_vendedores, top_compradores = ranking_participantes(conexion)
        return render_template(
            "estadisticas.html",
            activo="estadisticas",
            desglose=desglose_marcas(conexion),
            mensual=resumen_mensual(conexion),
            top_vendedores=top_vendedores,
            top_compradores=top_compradores,
            perfil=perfil_actual(),
        )

    @app.get("/aprobaciones")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def aprobaciones():
        conexion = conexion_actual()
        vista = request.args.get("vista") or "pendientes"
        return render_template(
            "aprobaciones.html",
            activo="aprobaciones",
            vista=vista,
            cambios=cambios_pendientes_detalle(conexion) if vista == "pendientes" else [],
            historial=historial_cambios_resueltos(conexion) if vista == "historial" else [],
            perfil=perfil_actual(),
        )

    @app.post("/aprobaciones/<int:cambio_id>")
    @requiere_sesion
    @requiere_rol("administrador", "operador")
    def resolver_cambio(cambio_id: int):
        conexion = conexion_actual()
        perfil = perfil_actual()
        decision = request.form.get("decision")
        motivo = request.form.get("motivo") or None
        try:
            resolver_cambio_pendiente(conexion, cambio_id, decision, perfil["id"], motivo)
        except ErrorResolverCambio as exc:
            print(f"No se pudo resolver el cambio {cambio_id}: {exc}")
            return render_template(
                "error_simple.html", titulo="No se pudo resolver el cambio", mensaje=str(exc)
            ), 400
        except Exception as exc:
            return _error_seguro("No se pudo resolver el cambio.", exc)
        return redirect(url_for("aprobaciones"))

    @app.get("/panel")
    @requiere_sesion
    @requiere_rol("administrador")
    def panel():
        conexion = conexion_actual()
        usuarios = listar_usuarios(conexion)
        return render_template(
            "panel.html", activo="panel", usuarios=usuarios, perfil=perfil_actual(), error=None
        )

    @app.get("/panel/respaldo.zip")
    @requiere_sesion
    @requiere_rol("administrador")
    def respaldo_completo():
        conexion = conexion_actual()
        datos = exportar_todo_para_respaldo(conexion)
        ahora = datetime.now(timezone.utc)
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
            for tabla, filas in datos.items():
                zf.writestr(f"{tabla}.json", dict_a_json(filas))
            zf.writestr(
                "manifiesto.json",
                dict_a_json({
                    "generado_en": ahora,
                    "generado_por": perfil_actual()["email"],
                    "cantidades": {tabla: len(filas) for tabla, filas in datos.items()},
                }),
            )
        buffer.seek(0)
        nombre = f"respaldo_piedra_alta_{ahora.strftime('%Y%m%d_%H%M%S')}.zip"
        return send_file(buffer, as_attachment=True, download_name=nombre, mimetype="application/zip")

    @app.post("/panel/usuarios")
    @requiere_sesion
    @requiere_rol("administrador")
    def crear_usuario_ruta():
        conexion = conexion_actual()
        email = request.form.get("email", "").strip().lower()
        nombre = request.form.get("nombre", "").strip()
        rol = request.form.get("rol") or "consulta"
        contrasena = request.form.get("contrasena", "")
        error = None
        if not email or not nombre:
            error = "Completá el email y el nombre."
        elif len(contrasena) < 6:
            error = "La contraseña debe tener al menos 6 caracteres."
        elif rol not in ("administrador", "operador", "consulta"):
            error = "Rol inválido."
        elif obtener_usuario_por_email(conexion, email):
            error = "Ya existe una cuenta con ese email."
        if error:
            return render_template(
                "panel.html", activo="panel", usuarios=listar_usuarios(conexion), perfil=perfil_actual(), error=error
            ), 400
        crear_usuario(conexion, email, encriptar_contrasena(contrasena), nombre, rol, True)
        return redirect(url_for("panel"))

    @app.post("/panel/usuarios/<usuario_id>")
    @requiere_sesion
    @requiere_rol("administrador")
    def actualizar_usuario(usuario_id: str):
        conexion = conexion_actual()
        if not obtener_usuario_por_id(conexion, usuario_id):
            return _registro_no_encontrado("No se encontró ese usuario.")
        campos = {}
        if request.form.get("rol"):
            campos["rol"] = request.form["rol"]
        if "activo" in request.form:
            campos["activo"] = request.form["activo"] == "1"
        nueva_contrasena = request.form.get("nueva_contrasena", "")
        if nueva_contrasena:
            if len(nueva_contrasena) < 6:
                return _error_seguro("La contraseña debe tener al menos 6 caracteres.", "contraseña corta")
            campos["password_hash"] = encriptar_contrasena(nueva_contrasena)
        if campos:
            actualizar_usuario_campos(conexion, usuario_id, campos)
        return redirect(url_for("panel"))

    return app


if __name__ == "__main__":       # pragma: no cover
    crear_app().run(debug=True)
