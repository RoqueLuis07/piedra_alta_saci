"""Utilidades compartidas por las rutas: errores, fechas, Excel, PDF y archivo."""

from __future__ import annotations

import base64
import io
import re
import tempfile
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pypdfium2 as pdfium
from flask import render_template, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

from marcas.pdf.guia import analizar_venta, completar_venta, extraer_encabezado
from marcas.servidor.auth import perfil_actual
from marcas.servidor.consultas import guardar_documento_emitido

_CODIGO_VALIDO = re.compile(r"^[A-Za-z0-9._-]+$")


def _error_seguro(mensaje: str, exc: Exception):
    """El detalle real queda en los logs del servidor, nunca en la respuesta:
    puede traer nombres de columnas o de tablas que no hace falta mostrar."""
    print(f"{mensaje}: {exc}")
    return render_template("error_simple.html", titulo="No se pudo completar", mensaje=mensaje), 400


def _registro_no_encontrado(mensaje: str):
    return render_template("error_simple.html", titulo="No encontrado", mensaje=mensaje), 404


def _fecha_sin_tz(valor):
    """openpyxl no admite datetimes con huso horario (falla al guardar el
    .xlsx) -- ``creado_en`` es timestamptz, así que hay que sacarle el tzinfo
    antes de escribirlo en una celda."""
    if isinstance(valor, datetime) and valor.tzinfo is not None:
        return valor.replace(tzinfo=None)
    return valor


_ZONA_HORARIA_LOCAL = ZoneInfo("America/Asuncion")


def _formatear_fecha_hora(valor):
    """Fecha y hora en un formato que cualquiera puede leer (DD/MM/AAAA
    HH:MM, en hora de Paraguay) -- nunca la marca de tiempo cruda que
    guarda la base (con microsegundos y huso horario UTC), que es jerga de
    base de datos y no algo pensado para mostrarse en pantalla."""
    if not valor:
        return "—"
    if isinstance(valor, str):
        try:
            valor = datetime.fromisoformat(valor)
        except ValueError:
            return valor
    if valor.tzinfo is not None:
        valor = valor.astimezone(_ZONA_HORARIA_LOCAL)
    return valor.strftime("%d/%m/%Y %H:%M")


# Nombre de columna -> cómo se llama ese dato en el resto de la interfaz --
# para que el detalle de un cambio en Aprobaciones diga "Vendedor" en vez
# de "vendedor_nombre" (jerga de base de datos que no tiene por qué
# conocer quien sólo entra a aprobar o rechazar un cambio).
_ETIQUETAS_CAMPO = {
    "numero_guia": "N.º de guía", "fecha": "Fecha",
    "vendedor_nombre": "Vendedor", "vendedor_documento": "CI / RUC del vendedor",
    "vendedor_establecimiento": "Establecimiento del vendedor",
    "vendedor_establecimiento_codigo": "Código de establecimiento del vendedor",
    "comprador_nombre": "Comprador", "comprador_documento": "CI / RUC del comprador",
    "cantidad_animales": "Cantidad de animales", "categoria_animales": "Categoría",
    "categoria_animales_original": "Categoría (original del formulario)",
    "tipo_formulario": "Tipo de formulario", "revisar": "Notas de revisar",
    "tipo_operacion": "Tipo de operación",
    "codigo": "Código", "tipo": "Tipo", "descripcion": "Descripción",
    "estado": "Estado", "observaciones": "Observaciones", "vence_en": "Vencimiento",
    "telefono": "Teléfono", "localidad": "Localidad", "departamento": "Departamento",
    "establecimiento": "Establecimiento", "establecimiento_codigo": "Código de establecimiento",
}


def _etiqueta_campo(campo: str) -> str:
    return _ETIQUETAS_CAMPO.get(campo, campo)


def _libro_excel(nombre_hoja: str, encabezados: list[str], filas: list[list]) -> io.BytesIO:
    """Arma un .xlsx prolijo para presentar como informe -- encabezado en
    negrita con fondo verde, congelado arriba, columnas con ancho ajustado
    al contenido y fechas con formato real (no todo como texto plano, como
    quedaba en el CSV viejo)."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre_hoja[:31]
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4D34")
        celda.alignment = Alignment(vertical="center")
    hoja.freeze_panes = "A2"
    for fila in filas:
        hoja.append(fila)
    for fila_celdas in hoja.iter_rows(min_row=2):
        for celda in fila_celdas:
            if isinstance(celda.value, datetime):
                celda.number_format = "DD/MM/YYYY HH:MM"
            elif isinstance(celda.value, date):
                celda.number_format = "DD/MM/YYYY"
    for indice, encabezado in enumerate(encabezados):
        valores = [encabezado] + [fila[indice] for fila in filas]
        ancho = max((len(str(v)) for v in valores if v is not None), default=0)
        hoja.column_dimensions[get_column_letter(indice + 1)].width = min(max(ancho + 2, 10), 42)
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer




_CAMPOS_MINIMOS_GUIA = ("numero_guia", "vendedor_nombre", "comprador_nombre", "cantidad_animales")


# "vendedor_nombre" no cuenta para una Venta: el formulario lo precarga
# siempre con los datos fijos de Piedra Alta (es quien vende), así que
# viene con algo cargado aunque la persona no haya tocado nada -- no sirve
# como señal de que el formulario tiene información real.
_CAMPOS_MINIMOS_VENTA = ("numero_guia", "comprador_nombre", "cantidad_animales")


def _operacion_vacia(campos: dict, minimos: tuple[str, ...] = _CAMPOS_MINIMOS_GUIA) -> bool:
    """True si ninguno de los campos que realmente identifican una guía o
    venta tiene algo cargado -- antes se podía crear un registro real
    presionando "Guardar" sin llenar nada (encontrado en la ronda de QA),
    dejando una guía o venta fantasma en la base sin que nadie se entere."""
    return all(not campos.get(c) for c in minimos)


def _preparar_respaldo(archivos) -> tuple[bytes, str]:
    """Arma el documento de respaldo permanente de una Compra a partir de lo
    subido en "Nueva guía": si es un único PDF, se guarda tal cual (es el
    caso de un escaneo ya armado en PDF); si son una o varias imágenes
    (fotos o páginas escaneadas sueltas), se combinan en un solo PDF de
    varias páginas, en el orden en que se subieron -- así el respaldo
    siempre queda como un solo documento, sea cual sea el formato de origen,
    y se puede ver/descargar igual que cualquier otro PDF de la operación.

    Deliberadamente no mezcla un PDF con imágenes en la misma carga (¿en qué
    orden entrarían las páginas del PDF respecto de las fotos?) -- levanta
    ValueError para que quien llama lo muestre como un error claro, en vez
    de que ``Image.open`` reviente más abajo con un traceback críptico."""
    if any(a.mimetype == "application/pdf" or a.filename.lower().endswith(".pdf") for a in archivos):
        if len(archivos) > 1:
            raise ValueError("no se puede combinar un PDF con otros archivos en la misma carga")
        archivo = archivos[0]
        return archivo.read(), archivo.filename

    # exif_transpose: una foto sacada con el celular en vertical trae los
    # píxeles en horizontal más una etiqueta EXIF de rotación -- el
    # navegador la respeta al mostrar la vista previa, pero PIL no la
    # aplica sola al abrir el archivo. Sin esto, el respaldo guardado
    # quedaría rotado aunque la vista previa que vio el operador estuviera
    # derecha.
    imagenes = [
        ImageOps.exif_transpose(Image.open(archivo.stream)).convert("RGB") for archivo in archivos
    ]
    buffer = io.BytesIO()
    imagenes[0].save(buffer, format="PDF", save_all=True, append_images=imagenes[1:])
    return buffer.getvalue(), "respaldo.pdf"






def _entregar_documento(conexion, operacion: dict, tipo: str, marcas: list[dict], contenido_pdf: bytes):
    """Cola común de generar guía / generar venta: archiva el PDF (fecha +
    datos de la operación + marcas) y lo manda a descargar.
    Si no se puede archivar, no se entrega: así nunca sale un documento que no
    quede guardado."""
    prefijo = "venta" if tipo == "venta" else "guia"
    nombre = f"{prefijo}_{operacion.get('numero_guia') or operacion['id']}.pdf"
    perfil = perfil_actual()
    try:
        guardar_documento_emitido(
            conexion, operacion, tipo, marcas, contenido_pdf, nombre, perfil["id"] if perfil else None
        )
    except Exception as exc:
        # Sin el rollback, Flask haría commit de la operación y las marcas
        # recién creadas aunque el documento no haya quedado archivado.
        conexion.rollback()
        return _error_seguro("No se pudo guardar el documento en el archivo, así que no se generó.", exc)
    return send_file(
        io.BytesIO(contenido_pdf), as_attachment=True, download_name=nombre, mimetype="application/pdf",
    )


def _estampar(pdf_bytes: bytes, dominante: bytes | None, complementarias: list[bytes]) -> tuple[bytes, dict]:
    """Completa la guía (Rubro 2 y, si sobran marcas, el Anexo) en un
    directorio temporal. Devuelve el PDF resultante y lo que informó
    ``completar_venta`` (páginas tocadas, cuántas entraron en cada lugar)."""
    with tempfile.TemporaryDirectory() as tmp:
        carpeta = Path(tmp)
        entrada = carpeta / "entrada.pdf"
        entrada.write_bytes(pdf_bytes)
        ruta_dominante = None
        if dominante:
            ruta_dominante = carpeta / "dominante.png"
            ruta_dominante.write_bytes(dominante)
        rutas = []
        for i, contenido in enumerate(complementarias):
            destino = carpeta / f"marca_{i:03d}.png"
            destino.write_bytes(contenido)
            rutas.append(destino)
        salida = carpeta / "salida.pdf"
        info = completar_venta(entrada, ruta_dominante, rutas, salida)
        return salida.read_bytes(), info


def _renderizar_paginas(pdf_bytes: bytes, indices: list[int], escala: float = 1.3) -> list[str]:
    """Las páginas pedidas del PDF como PNG en base64 (para la vista previa)."""
    resultado: list[str] = []
    doc = pdfium.PdfDocument(pdf_bytes)
    try:
        for i in indices:
            if i >= len(doc):
                continue
            buffer = io.BytesIO()
            doc[i].render(scale=escala).to_pil().save(buffer, format="PNG")
            resultado.append(base64.b64encode(buffer.getvalue()).decode("ascii"))
    finally:
        doc.close()
    return resultado


def _analizar_pdf(pdf_bytes: bytes) -> dict:
    """Cupo libre y datos leídos del encabezado del PDF, o ``error`` si no
    parece una Guía de Traslado de SENACSA."""
    with tempfile.TemporaryDirectory() as tmp:
        entrada = Path(tmp) / "entrada.pdf"
        entrada.write_bytes(pdf_bytes)
        try:
            cupo = analizar_venta(entrada)
        except Exception as exc:
            return {"cupo": None, "encabezado": None, "error": str(exc)}
        try:
            encabezado = extraer_encabezado(entrada)
        except Exception:
            encabezado = None
    return {"cupo": cupo, "encabezado": encabezado}
